from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SRC = Path(r"D:\AI Sign Language Translator\AI Sign translator Book report.docx")
OUT = Path(r"D:\AI Sign Language Translator\AI Sign translator Book report_detailed.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.size = Pt(16 if level == 1 else 14 if level == 2 else 12)
    return paragraph


def add_para(doc, text="", bold_label=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    if bold_label:
        label = paragraph.add_run(bold_label)
        label.bold = True
        label.font.name = "Times New Roman"
        label.font.size = Pt(12)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return paragraph


def add_list(doc, items):
    for item in items:
        paragraph = doc.add_paragraph("- " + item)
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph(f"{index}. {item}")
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)


def add_table(doc, caption, headers, rows):
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(3)
    run = cap.add_run(caption)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
        set_cell_shading(table.rows[0].cells[i], "D9EAF7")

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))

    doc.add_paragraph()
    return table


def add_fig(doc, number, description):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(f"Figure {number}: [INSERT IMAGE: {description}]")
    run.bold = True
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(31, 78, 121)
    return paragraph


def configure_doc(doc):
    if "Normal" in doc.styles:
        normal = doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def build():
    copyfile(SRC, OUT)
    doc = Document(OUT)
    configure_doc(doc)
    doc.add_page_break()

    add_heading(doc, "3.5 Dataset Preparation and Label Arrangement", 2)
    add_para(doc, "The dataset preparation stage was redesigned to support the final scope of the application, which focuses on fingerspelling recognition rather than broad sentence-level sign recognition. In the earlier system version, the dataset included letter images, letter videos, word videos, and sentence videos. This made the training task too broad for the available dataset size and caused unstable prediction behavior. The refined version limits the training domain to alphabet-level ASL fingerspelling, using both image samples and letter video samples. This means that image samples of letter A and video samples of letter A are treated as the same class, and the same rule is applied from B to Z.")
    add_para(doc, "The purpose of this design choice is to make the model learn a more consistent relationship between hand pose, finger-joint arrangement, palm orientation, and the predicted letter. Instead of asking the model to recognize full words or full sentences directly, the application builds words letter by letter. This is better for a real-time translator because each second of signing produces one candidate letter, and the completed set of letters can be spoken when the hand leaves the frame.")
    add_para(doc, "The dataset is stored in two main groups. The first group contains still images located inside data/images, arranged by letter folders. The second group contains short video clips of letter signs. Each image or video is indexed with its alphabet label. During training, both sample types are loaded through the same dataset class so that the model receives a unified class target.")
    add_fig(doc, "3.1", "Dataset folder structure showing data/images/A-Z folders and letter video folders")
    add_table(doc, "Table 3.1: Final Dataset Categories Used for Training", ["Dataset Type", "Included?", "Purpose in Training", "Reason"], [
        ["Letter images", "Yes", "Provide many static hand-pose examples for each alphabet class.", "Improves learning of finger arrangement and palm shape."],
        ["Letter videos", "Yes", "Provide temporal movement and webcam-like samples for each alphabet class.", "Improves real-time prediction behavior."],
        ["Word videos", "No", "Removed from the final alphabet model.", "Prevents word-level samples from confusing single-letter output."],
        ["Sentence videos", "No", "Removed from the final alphabet model.", "Keeps the model focused on one-second letter prediction."],
    ])
    add_para(doc, "Before training, the dataset indexer scans the available folders, validates file paths, checks that labels are known alphabet classes, and removes unusable samples. In particular, image samples that do not contain real MediaPipe hand landmarks are dropped. This prevents the model from learning from corrupted or ambiguous examples. The removal of samples without hands is important because a no-hand image should not be treated as a valid letter training example.")

    add_heading(doc, "3.6 Landmark Extraction and Hand Feature Representation", 2)
    add_para(doc, "The system uses hand landmarks as a second input stream alongside the CNN image input. The landmark stream is important because many ASL letters are defined by finger configuration rather than only by raw image appearance. For example, the difference between letters such as A and E, or Q and W, may depend on the relative position of the thumb, index finger, palm center, and finger joints. A CNN can learn visual appearance, but landmarks give the model direct geometric information about the hand.")
    add_para(doc, "The landmark extractor is based on MediaPipe hand detection. It identifies the wrist, thumb joints, index finger joints, middle finger joints, ring finger joints, and pinky joints. These points are converted into a feature vector that includes normalized x-y-z landmark coordinates, hand presence confidence, bounding-box information, and geometric measurements. Contour-based detection is intentionally not used for prediction because earlier testing showed that contour detection responded to skin-colored regions and background noise instead of reliable hand structure.")
    add_fig(doc, "3.2", "MediaPipe hand landmarks drawn on a live webcam frame")
    add_para(doc, "The hand landmark process works in the following sequence. First, the frame is passed to MediaPipe. Second, if a hand is found, the 21 hand landmarks are collected. Third, the coordinates are normalized so that the model receives stable measurements even when the hand is closer or farther from the camera. Fourth, the same landmarks are used to create a tight crop around the hand. Finally, both the cropped hand image and landmark vector are sent to the model.")
    add_numbered(doc, [
        "Capture a frame from the webcam or read a training image/video frame.",
        "Run true MediaPipe hand landmark detection on the frame.",
        "Reject the frame if no reliable hand is detected.",
        "Generate a landmark feature vector from wrist, palm, thumb, index, middle, ring, and pinky joints.",
        "Crop the image around the detected hand region using the landmark bounding box.",
        "Resize and normalize the crop for the CNN branch.",
        "Pass both CNN image features and landmark features into the hybrid model.",
    ])
    add_table(doc, "Table 3.2: Landmark-Based Features Used by the Model", ["Feature Group", "Description", "Benefit"], [
        ["Wrist and palm region", "Used as the base reference for hand position and orientation.", "Helps the model understand palm direction and hand center."],
        ["Thumb joints", "Tracks thumb tip and intermediate thumb joints.", "Important for letters such as A, E, M, N, S, T, and Y."],
        ["Index finger joints", "Tracks index base, middle, upper, and tip landmarks.", "Important for letters such as D, G, H, I, L, Q, U, V, W, and X."],
        ["Middle finger joints", "Tracks middle finger position and bending.", "Helps distinguish U, V, W, and similar hand shapes."],
        ["Ring and pinky joints", "Tracks closed or open finger states.", "Important for letters B, F, I, W, and Y."],
        ["Hand presence score", "Stores whether a reliable hand exists in the frame.", "Supports no-hand detection and prevents false predictions."],
        ["Bounding box values", "Stores hand crop position and relative size.", "Supports tight crop consistency between training and prediction."],
    ])

    add_heading(doc, "3.7 Model Training Strategy", 2)
    add_para(doc, "The model was trained as a hybrid architecture because the final application requires both visual and structural understanding of a sign. The CNN branch learns local visual details such as palm shape, finger silhouette, lighting variation, and hand texture. The landmark branch learns geometric relationships between joints. The temporal branch learns how the sequence behaves across the one-second buffer. Together, these branches form a stronger alphabet classifier than a CNN-only model.")
    add_para(doc, "Training uses a combination of Connectionist Temporal Classification (CTC) loss and an auxiliary classification loss. CTC supports sequential predictions from frame sequences, while the auxiliary classifier helps the model learn a direct alphabet target. This is helpful because the final app expects one letter per second. Class weighting and weighted sampling are applied so that minority classes are not ignored. The sampler balances image and letter-video sources so that the model does not become biased toward only one data type.")
    add_fig(doc, "3.3", "Hybrid training flow from dataset loader to CNN branch, landmark branch, temporal module, CTC head, and classifier head")
    add_table(doc, "Table 3.3: Main Training Configuration", ["Parameter", "Value / Strategy", "Reason"], [
        ["Training classes", "A-Z alphabet letters only", "Keeps the model focused on fingerspelling output."],
        ["Image samples", "Included", "Improves static hand-shape recognition."],
        ["Letter video samples", "Included", "Improves real-time temporal behavior."],
        ["Word/sentence samples", "Excluded", "Avoids mixed objectives during alphabet prediction."],
        ["Landmarks", "Enabled", "Adds finger-joint and palm geometry to training."],
        ["Tight crop", "Enabled in both training and live prediction", "Keeps model input consistent."],
        ["Class weights", "Enabled", "Reduces dominance of frequent classes."],
        ["Weighted sampler", "Enabled", "Balances images and videos during training."],
        ["Validation metrics", "CER, exact match, precision, recall, F1", "Gives understandable model quality feedback."],
    ])
    add_para(doc, "During training, the validation process reports overall metrics as well as separate image and letter-video performance. This separation is important because a model may perform well on still images but fail on live-like video samples. The training process also saves checkpoints and tracks the best model based on validation quality, not just training loss. This avoids selecting an overfitted model that memorizes the training data but fails during webcam prediction.")

    add_heading(doc, "3.8 Real-Time Prediction Method", 2)
    add_para(doc, "The real-time prediction stage was designed around one-second decision windows. The webcam captures frames continuously, and the inference worker processes frames using the same preprocessing mechanism used during training. Each frame is mirrored for a natural user experience, landmarks are extracted, a tight crop is created, and the hybrid model predicts a candidate letter with confidence.")
    add_para(doc, "At the end of each one-second prediction window, the system collects all frame-level predictions from that window and chooses one final letter. The first decision rule is majority voting: the letter that appears most often within the one-second window is selected. If two or more letters appear with the same frequency, the system compares confidence values and chooses the candidate with the highest confidence. After the one-second result is produced, the temporary prediction cache for that second is cleared so that old predictions do not influence the next second.")
    add_fig(doc, "3.4", "One-second prediction window showing frame predictions, majority vote, confidence tie-break, and cache reset")
    add_numbered(doc, [
        "Frames are captured by the camera thread and sent to the frame queue.",
        "The inference worker reads frames, extracts landmarks, crops the hand, and predicts candidate letters.",
        "Predictions from the current one-second window are stored temporarily.",
        "At the one-second boundary, the most frequent letter is selected.",
        "If there is a tie, the prediction with the highest confidence is selected.",
        "The selected letter is appended to the current letter set.",
        "The one-second prediction cache is cleared before the next window begins.",
    ])
    add_para(doc, "This method is stronger than using a single frame because it reduces noise from accidental hand movement, lighting flicker, momentary detection errors, and unstable poses. It also gives the user a predictable rhythm: one letter is produced approximately every second while the hand remains visible.")

    add_heading(doc, "3.9 No-Hand Detection and Word Completion Logic", 2)
    add_para(doc, "No-hand detection is one of the most important parts of the final system. Without it, the application may continue predicting letters even when the user removes their hand from the camera. This creates false letters and prevents the system from knowing when to speak the completed word. The refined design uses real MediaPipe hand presence instead of color contour detection. A frame is treated as no-hand when MediaPipe fails to produce reliable hand landmarks or when the hand confidence is below the selected threshold.")
    add_para(doc, "The application does not speak immediately after one no-hand frame because a single missing detection may happen because of motion blur, camera exposure, or partial hand occlusion. Instead, the system counts consecutive no-hand frames. When the number of consecutive no-hand frames reaches the configured threshold, the current letter set is considered complete. The completed set is then added to the completed words list and sent to the speech worker for text-to-speech output.")
    add_fig(doc, "3.5", "No-hand detection flow showing hand present state, no-hand counter, word flush, and speech output")
    add_table(doc, "Table 3.4: No-Hand Completion Rules", ["Condition", "System Action", "Reason"], [
        ["Hand detected", "Continue collecting one-second letter predictions.", "The user is still signing."],
        ["One or few no-hand frames", "Wait and do not speak yet.", "Avoids false completion from temporary tracking loss."],
        ["No-hand threshold reached", "Flush current letter set and speak it.", "Signals that the user finished the word."],
        ["No letters collected", "Do not speak.", "Prevents empty speech output."],
        ["User resets", "Clear prediction cache and current letter set.", "Allows recovery from mistakes."],
    ])

    add_heading(doc, "3.10 Threaded System Architecture", 2)
    add_para(doc, "The real-time app uses a threaded architecture so that camera capture, model inference, display updates, and speech output do not block each other. This is important because webcam applications can become unstable if every operation runs inside one loop. For example, text-to-speech can pause the program while speaking, and model inference can be slower on CPU. By separating responsibilities into workers and queues, the app remains more responsive.")
    add_fig(doc, "3.6", "Threaded runtime architecture: camera thread, frame queue, inference worker, prediction queue, UI display loop, and speech worker")
    add_table(doc, "Table 3.5: Real-Time Worker Responsibilities", ["Worker / Loop", "Responsibility", "Input", "Output"], [
        ["Camera thread", "Captures frames from webcam and mirrors the live feed.", "Webcam device", "Frame queue"],
        ["Inference worker", "Runs landmarks, tight crop, CNN-landmark model, and prediction logic.", "Frame queue", "Prediction queue"],
        ["UI/display loop", "Draws live feed, landmarks, tabs, current predictions, and completed sets.", "Latest frame and predictions", "OpenCV window"],
        ["Speech worker", "Speaks completed letter sets without blocking camera or UI.", "Speech queue", "Audio output"],
        ["Microphone worker", "Handles speech-to-sign voice recording and recognition.", "Microphone input", "Recognized text"],
    ])
    add_para(doc, "This architecture supports both sign-to-speech and speech-to-sign in one native Python application. The user can switch between tabs without opening a browser. In sign-to-speech mode, the live feed appears first and the prediction content appears below. In speech-to-sign mode, microphone controls and recognized text appear above, while the corresponding sign images are displayed below.")

    add_heading(doc, "CHAPTER FOUR", 1)
    add_heading(doc, "DESIGN AND IMPLEMENTATION", 1)
    add_heading(doc, "4.1 Introduction", 2)
    add_para(doc, "This chapter explains how the proposed AI Sign Language Translator was implemented. It focuses on the software modules, model design, data processing pipeline, real-time camera interface, speech-to-sign feature, and system testing structure. The implementation was designed to support two related services in one application: sign-to-speech and speech-to-sign. The sign-to-speech service recognizes ASL fingerspelling from a webcam and speaks completed letter sets. The speech-to-sign service listens to spoken input, converts it into text, and displays corresponding ASL letter images from the dataset.")
    add_para(doc, "The final implementation avoids browser dependency and runs as a native Python/OpenCV interface. This decision was made because the user interface must handle webcam frames, landmarks, microphone recording, and text-to-speech in real time. A native interface also makes it easier to inspect the live feed, draw landmarks, display sign images, and test prediction speed without Streamlit or browser refresh issues.")

    add_heading(doc, "4.2 System Overview", 2)
    add_para(doc, "The system is divided into five major layers: dataset layer, preprocessing layer, model layer, inference layer, and user interface layer. The dataset layer loads images and letter videos. The preprocessing layer extracts MediaPipe landmarks and creates tight hand crops. The model layer predicts alphabet classes using CNN and landmark features. The inference layer converts frame predictions into one-second letter decisions. The user interface layer shows live feed, prediction history, completed sets, microphone status, recognized speech text, and the speech-to-sign sign sequence.")
    add_fig(doc, "4.1", "Overall system architecture showing dataset, preprocessing, model, inference, UI, and speech output layers")
    add_table(doc, "Table 4.1: System Modules and Their Roles", ["Module", "Main Files", "Role"], [
        ["Dataset handling", "data/real_dataset.py", "Indexes letter images and videos, loads frames, labels samples, and returns tensors."],
        ["Landmark extraction", "utils/hand_landmarks.py", "Uses MediaPipe to extract hand landmarks and create tight crops."],
        ["Model architecture", "models/asl_ctc_transformer.py, models/temporal_conv.py", "Combines CNN image features, landmark features, temporal modeling, and prediction heads."],
        ["Training", "train_ctc.py", "Trains alphabet model, validates metrics, saves checkpoints, and reports performance."],
        ["Real-time app", "realtime/inference_ctc.py", "Runs sign-to-speech and speech-to-sign in native OpenCV interface."],
    ])

    add_heading(doc, "4.3 Data Loading Implementation", 2)
    add_para(doc, "The data loader was implemented to ensure that images and letter videos are treated as samples of the same alphabet classification problem. The loader reads labels from folder names and file paths, converts labels into class IDs, and returns both visual tensors and landmark tensors. For image samples, the image is repeated or converted into a frame sequence so that the model receives a consistent temporal shape. For video samples, frames are sampled across the clip and normalized to the expected length.")
    add_para(doc, "A key implementation detail is landmark caching. Extracting MediaPipe landmarks repeatedly during training can be slow, especially on CPU. To reduce repeated computation, extracted landmark features are saved and reused when possible. This makes later training runs faster and more stable. The dataset also filters out image samples where true hand landmarks cannot be detected, because those samples would introduce noise into the model.")
    add_fig(doc, "4.2", "Example of a training image before landmark extraction and after hand landmark detection")
    add_fig(doc, "4.3", "Example of tight hand crop generated from landmark bounding box")

    add_heading(doc, "4.4 Preprocessing Implementation", 2)
    add_para(doc, "The preprocessing stage is responsible for converting raw camera frames or dataset images into model-ready input. It performs frame conversion, MediaPipe processing, hand crop generation, resizing, normalization, and landmark feature extraction. The same preprocessing method is applied during training and live prediction so that the model does not see one type of image during training and a different type during inference.")
    add_para(doc, "Tight hand cropping was added because full webcam frames contain background objects, face regions, clothing, wall color, and lighting variation. These unrelated details can confuse the CNN. By cropping around the detected hand, the model receives a clearer view of the sign. However, the crop must be used consistently. If training uses full images but live prediction uses tight crops, the model may fail. Therefore, the implementation uses landmark-based tight crops in both training and live prediction.")
    add_table(doc, "Table 4.2: Preprocessing Consistency Rules", ["Rule", "Training", "Live Prediction"], [
        ["Use MediaPipe hand landmarks", "Yes", "Yes"],
        ["Use contour fallback", "No", "No"],
        ["Use tight hand crop", "Yes", "Yes"],
        ["Use same image size", "Yes", "Yes"],
        ["Use landmark feature vector", "Yes", "Yes"],
        ["Reject no-hand frames", "Yes for invalid samples", "Yes for prediction and word completion"],
    ])

    add_heading(doc, "4.5 Hybrid CNN-Landmark Model Implementation", 2)
    add_para(doc, "The model combines image-based and landmark-based learning. The CNN branch receives cropped hand frames and extracts visual features. The landmark branch receives numerical hand features and learns the structural relationship between the wrist, palm, thumb, and finger joints. The outputs from both branches are fused and passed through a temporal module. The temporal module considers the frame sequence and produces time-aware features. Finally, the model has a CTC prediction head and an auxiliary classifier head.")
    add_fig(doc, "4.4", "Detailed model architecture showing CNN encoder, landmark encoder, fusion layer, temporal convolution, CTC head, and classifier head")
    add_table(doc, "Table 4.3: Model Component Details", ["Component", "Input", "Output", "Purpose"], [
        ["CNN encoder", "Cropped hand image frames", "Visual feature vectors", "Learns palm shape, finger silhouette, and image texture."],
        ["Landmark encoder", "Hand landmark feature vector", "Geometric feature vectors", "Learns finger-joint arrangement and hand pose."],
        ["Fusion layer", "CNN + landmark features", "Combined representation", "Allows visual and structural information to support each other."],
        ["Temporal module", "Frame sequence features", "Time-aware features", "Reduces frame-level noise and captures short motion."],
        ["CTC head", "Temporal features", "Character logits over time", "Supports sequence-style training and decoding."],
        ["Classifier head", "Aggregated features", "Single alphabet class logits", "Supports one-letter-per-second prediction."],
    ])
    add_para(doc, "The model is intentionally not trained to output words or sentences directly. Instead, it outputs letters. This design fits the real-time fingerspelling workflow, where the user signs letters one after another and the system builds the word from the predicted letter sequence.")

    add_heading(doc, "4.6 Training Implementation", 2)
    add_para(doc, "Training is executed through train_ctc.py. The script builds the dataset, splits samples into training and validation sets, creates a weighted sampler, initializes the model, and performs epoch-based optimization. The training process reports batch loss, train loss, validation loss, character error rate, exact match, precision, recall, and F1 score. It also reports performance separately for image samples and letter-video samples, which helps reveal whether the model is only learning still images or also learning video-based signs.")
    add_para(doc, "The training script saves checkpoint files for each epoch and saves the best model separately as checkpoints/best_model.pt. The best model is selected based on validation performance. This is important because the lowest training loss is not always the best real-world model. In earlier testing, the training loss could become low while live prediction still failed because the model overfitted to the training examples.")
    add_fig(doc, "4.5", "Training log screenshot showing validation loss, CER, exact match, precision, recall, and F1 score")
    add_table(doc, "Table 4.4: Training Metrics Explained", ["Metric", "Meaning", "Why It Matters"], [
        ["Validation loss", "Loss measured on unseen validation samples.", "Shows whether the model generalizes beyond training samples."],
        ["CER", "Character Error Rate.", "Measures how many character-level mistakes occur. Lower is better."],
        ["Exact match", "Percentage of samples where predicted text exactly matches target.", "Easy to interpret as accuracy for letter samples."],
        ["Precision", "How many predicted letters are correct.", "Shows false prediction behavior."],
        ["Recall", "How many true letters are found correctly.", "Shows missed-class behavior."],
        ["F1 score", "Balance of precision and recall.", "Useful when class distribution is not perfectly balanced."],
    ])

    add_heading(doc, "4.7 Sign-to-Speech Interface Implementation", 2)
    add_para(doc, "The sign-to-speech interface opens the webcam, mirrors the live feed, detects the hand, draws hand landmarks, and displays prediction information in a content section. The live feed is kept clean and is not overloaded with long text. Prediction results such as all predictions and completed sets are displayed in the content area so that the user can understand the system state without covering the hand region.")
    add_para(doc, "The interface displays the latest one-second prediction, the current letter set, all predictions in sequence, and completed sets. When no hand is detected for the configured number of frames, the current set is moved to completed sets and spoken through text-to-speech. This makes the system work more naturally because the user can spell a word and then remove the hand to trigger speech output.")
    add_fig(doc, "4.6", "Sign-to-speech screen showing live feed at top and prediction content section below")
    add_fig(doc, "4.7", "Sign-to-speech content area showing All predictions: L -> S -> Z and Completed sets: LSZ")

    add_heading(doc, "4.8 Speech-to-Sign Interface Implementation", 2)
    add_para(doc, "The speech-to-sign interface works as the inverse of sign-to-speech. In this mode, the user speaks through the microphone. The app captures the voice input, converts it to text, and displays corresponding sign images for each letter in the recognized text. The images are selected from data/images using the folder that matches each letter. For example, if the recognized text contains the letter A, the app selects an image from data/images/A.")
    add_para(doc, "The interface displays microphone status, recognized text, and a sign output screen. The sign output does not show one image at a time only. Instead, it displays the full sequence of signs at once, grouped according to words where possible. This makes the output easier to understand because the user can see the spelled form of the spoken phrase rather than waiting for a single animated sign at a time.")
    add_fig(doc, "4.8", "Speech-to-sign screen showing microphone recording status and recognized text above the sign output area")
    add_fig(doc, "4.9", "Speech-to-sign sign output showing all selected letter images arranged in sequence")
    add_table(doc, "Table 4.5: Speech-to-Sign Processing Steps", ["Step", "Description"], [
        ["Microphone selection", "The app allows the user to list and choose microphone devices by index."],
        ["Voice capture", "Audio is captured from the selected microphone."],
        ["Speech recognition", "Recognized speech is converted into text."],
        ["Text cleaning", "Text is converted to uppercase letters and unsupported symbols are ignored."],
        ["Image selection", "For each letter, one matching sign image is selected from data/images."],
        ["Output display", "All signs are displayed in sequence in the speech-to-sign output screen."],
    ])

    add_heading(doc, "4.9 Microphone Handling and Voice Recognition", 2)
    add_para(doc, "Microphone support was implemented with practical debugging options because voice input can fail for several reasons: the wrong microphone index may be selected, the microphone may be too quiet, background noise may be high, or the speech recognizer may time out before detecting a phrase. The app therefore includes microphone listing, microphone index selection, timeout handling, and status messages that explain whether the microphone is listening, timed out, heard unclear speech, or returned recognized text.")
    add_fig(doc, "4.10", "Microphone device list and selected mic-index test output")
    add_para(doc, "The app also separates microphone logic from the camera loop so that listening does not freeze the live camera. This design is especially useful because speech recognition may take several seconds, while the sign-to-speech mode still needs responsive display behavior.")

    add_heading(doc, "4.10 Queue and Worker Implementation", 2)
    add_para(doc, "The queue-based architecture prevents slow operations from blocking each other. Camera capture, inference, display, speech output, and microphone recognition each have separate responsibilities. Frame queues move captured frames to inference. Prediction queues move model results to the display loop. Speech queues move completed sets to the text-to-speech worker.")
    add_table(doc, "Table 4.6: Queue-Based Runtime Flow", ["Queue / Worker", "Producer", "Consumer", "Data Passed"], [
        ["Frame queue", "Camera thread", "Inference worker", "Mirrored webcam frames"],
        ["Prediction queue", "Inference worker", "Display loop", "Latest prediction, confidence, hand status, landmarks"],
        ["Speech queue", "No-hand completion logic", "Speech worker", "Completed letter sets"],
        ["Microphone result state", "Microphone worker", "Speech-to-sign display", "Recognized text and status messages"],
    ])
    add_fig(doc, "4.11", "Runtime queue diagram showing non-blocking camera, inference, display, and speech workers")

    add_heading(doc, "4.11 Implementation Challenges and Corrections", 2)
    add_para(doc, "Several issues were identified during development and corrected. The first major issue was blank prediction dominance, where the model produced empty output or repeated the same character. This was improved through training adjustments, checkpoint selection, and live decoding controls. The second issue was class bias, especially toward some letters such as S. This was addressed by checking dataset distribution, applying class weights, and improving sampling balance. The third issue was false prediction when no hand was visible. This was corrected using real hand landmark detection and a no-hand frame threshold.")
    add_para(doc, "Another important correction was the removal of word and sentence datasets from alphabet training. Including word and sentence samples made the model objective inconsistent with one-letter-per-second prediction. The final system treats images and letter videos of the same letter as one class. This gives the model a clearer target and supports the application goal of spelling words from alphabet predictions.")
    add_table(doc, "Table 4.7: Bugs Found and Fixes Applied", ["Problem Found", "Cause", "Fix Applied"], [
        ["Module import error in realtime script", "Project root was not on Python path.", "Added project-root path handling for realtime inference."],
        ["Repeated S/W/PP predictions", "Model/checkpoint bias and weak live gating.", "Improved checkpoint selection, class weighting, and no-hand gating."],
        ["No-hand still predicted letters", "Frame accepted even without reliable hand landmarks.", "Added MediaPipe-based no-hand threshold and word flush logic."],
        ["Contour detected wrong objects", "Skin/background color was treated as hand.", "Disabled contour fallback for prediction."],
        ["Training too broad", "Words, sentences, letters, and images mixed objectives.", "Restricted final model to alphabet images and letter videos."],
        ["Live feed not matching training", "Preprocessing differences between training and prediction.", "Applied landmarks and tight crop in both training and live inference."],
        ["Speech-to-sign output too limited", "Signs were not presented as complete sequence.", "Updated output to show all signs grouped by text."],
    ])

    add_heading(doc, "CHAPTER FIVE", 1)
    add_heading(doc, "TESTING, DISCUSSION, CONCLUSION AND RECOMMENDATIONS", 1)
    add_heading(doc, "5.1 Introduction", 2)
    add_para(doc, "This chapter presents the testing approach, system results, discussion of findings, limitations, conclusion, and recommendations. Testing was performed at different levels because an AI application can fail in many places: dataset loading, landmark extraction, model training, webcam inference, no-hand detection, speech output, microphone recognition, and user interface display. Testing each part separately helped identify the real source of errors instead of assuming that every problem came from the neural network.")

    add_heading(doc, "5.2 Testing Strategy", 2)
    add_para(doc, "The testing strategy followed a staged development workflow. Each layer of the project was validated before the next layer was trusted. Dataset validation came first, followed by landmark extraction testing, model forward-pass testing, training validation, real-time webcam testing, no-hand testing, and speech-to-sign testing. This staged approach reduced debugging complexity and made it easier to detect whether an error was caused by data, model architecture, training configuration, or live inference logic.")
    add_table(doc, "Table 5.1: System Testing Plan", ["Test Area", "Test Performed", "Expected Result"], [
        ["Dataset indexing", "Run dataset scan for images and letter videos.", "Only A-Z labels are indexed."],
        ["Landmark extraction", "Display landmarks on sample images and webcam frames.", "Hand joints appear on wrist, palm, and fingers."],
        ["Tight crop", "Compare full image and cropped hand input.", "Crop focuses on the hand region."],
        ["Model forward pass", "Pass dummy and real batches through model.", "No tensor shape errors or NaN values."],
        ["Training metrics", "Train and validate model.", "CER decreases and exact/F1 improve."],
        ["Live prediction", "Perform different signs in webcam.", "One letter result appears per second."],
        ["No-hand detection", "Remove hand from frame.", "System stops prediction and speaks completed set."],
        ["Speech-to-sign", "Speak a word into microphone.", "Text appears and matching sign images are displayed."],
    ])
    add_fig(doc, "5.1", "Testing screenshot showing model validation metrics after training")

    add_heading(doc, "5.3 Training Results", 2)
    add_para(doc, "The refined training process produced strong validation results after the system was limited to images and letter videos. The validation reports showed high exact match, precision, recall, and F1 score. In one training run, the validation exact match reached approximately 95.11%, with image performance around 97.22% and letter-video performance around 88.00%. These results indicate that the model learned the alphabet classes much better after the training objective was narrowed and the landmark/crop pipeline was made consistent.")
    add_para(doc, "The difference between image performance and letter-video performance is important. Images are usually cleaner and more centered than live camera frames, while videos contain motion, blur, background variation, and signer variation. Therefore, the letter-video score is a better signal for real-time performance. The result suggests that the model is usable, but additional real webcam data from the same environment could still improve prediction quality.")
    add_fig(doc, "5.2", "Training result screenshot showing Val Loss, CER, Exact, Precision, Recall, and F1")
    add_table(doc, "Table 5.2: Example Validation Result Summary", ["Metric", "Overall", "Images", "Letter Videos"], [
        ["CER", "0.0489", "0.0278", "0.1200"],
        ["Exact match", "0.9511", "0.9722", "0.8800"],
        ["Precision", "0.9511", "0.9722", "0.8800"],
        ["Recall", "0.9511", "0.9722", "0.8800"],
        ["F1 score", "0.9511", "0.9722", "0.8800"],
    ])

    add_heading(doc, "5.4 Real-Time Sign-to-Speech Discussion", 2)
    add_para(doc, "In real-time use, the system is expected to produce one letter per second while the hand remains visible. This design gives the user enough time to form each sign clearly. The application displays all predictions and the current completed sets so that the user can understand what the model is building. The no-hand logic acts as a word boundary. Once the user removes the hand, the system treats the current letter set as complete and sends it to the speech worker.")
    add_para(doc, "The live system depends heavily on hand detection quality. Poor lighting, fast movement, partial hand visibility, and unusual camera angles can reduce landmark accuracy. Since the model uses both landmarks and the cropped image, failed landmarks can affect both input streams. For this reason, the live feed includes landmark drawing so that the user can see whether the system is actually tracking the hand.")
    add_fig(doc, "5.3", "Final sign-to-speech prototype showing mirrored live feed with visible hand landmarks")

    add_heading(doc, "5.5 Speech-to-Sign Discussion", 2)
    add_para(doc, "The speech-to-sign feature extends the project beyond recognition by allowing communication in the opposite direction. A hearing user can speak into the microphone, and the application converts the recognized text into a visible ASL fingerspelling sequence. This makes the application more useful in two-way communication situations. The current implementation focuses on alphabet signs from the dataset rather than full animated sign language grammar.")
    add_para(doc, "Microphone recognition can be affected by background noise, microphone sensitivity, incorrect device index, and speech clarity. The app therefore includes microphone status messages so that the user can know whether the app is listening, timed out, heard unclear speech, or successfully recognized text. Future versions can improve this feature by adding offline speech recognition or a more robust speech model.")
    add_fig(doc, "5.4", "Final speech-to-sign prototype showing recognized text and full sign sequence output")

    add_heading(doc, "5.6 System Limitations", 2)
    add_para(doc, "Although the system achieved promising validation results, it still has limitations. The first limitation is dataset size. The model can only generalize well if it sees enough examples of different hands, lighting conditions, backgrounds, camera qualities, and signing styles. The second limitation is alphabet-only recognition. The final model intentionally excludes words and sentences, so it cannot directly understand full ASL grammar. The third limitation is real-time ambiguity. Some letters are visually similar, and even landmark-based features may struggle when the hand pose is unclear or the signer forms the sign loosely.")
    add_list(doc, [
        "The system is limited to ASL alphabet fingerspelling rather than complete ASL grammar.",
        "Performance may drop under poor lighting, motion blur, or unusual camera angle.",
        "Letters with similar hand shapes may still be confused if the signer posture is unclear.",
        "Speech-to-sign depends on microphone quality and speech recognition reliability.",
        "The dataset still needs more live webcam examples from multiple users for stronger generalization.",
    ])

    add_heading(doc, "5.7 Recommendations", 2)
    add_para(doc, "The first recommendation is to collect a custom live webcam dataset for each alphabet letter. The user should record multiple examples of A through Z under the same camera conditions used by the final application. This dataset should include different distances, slight rotations, different lighting levels, and natural hand movement. Such data would make the model better at real-world prediction because it would reduce the gap between training data and live webcam input.")
    add_para(doc, "The second recommendation is to build a guided data collection tool inside the project. The tool should open the camera, display the target letter, count down, record several seconds of signing, extract landmarks, save cropped frames, and store labels automatically. This would make dataset expansion easier and reduce labeling mistakes.")
    add_para(doc, "The third recommendation is to add a confusion report after validation. The confusion report should show which letters are mistaken for each other, such as A/E or Q/W. This would help decide which signs need more training samples. The fourth recommendation is to evaluate the system on users who were not included in the training data, because real generalization can only be confirmed with unseen signers.")
    add_fig(doc, "5.5", "Proposed future data collection screen showing target letter, countdown, camera preview, and recording status")

    add_heading(doc, "5.8 Conclusion", 2)
    add_para(doc, "This project developed an AI Sign Language Translator focused on ASL alphabet fingerspelling. The final system combines CNN image features, MediaPipe hand landmarks, tight hand cropping, temporal prediction, no-hand completion, text-to-speech, microphone input, and speech-to-sign display. The project evolved from a broad sign recognition system into a more focused and practical alphabet spelling system. This decision improved model reliability because the model was trained on a clear target: images and letter videos of A-Z classes.")
    add_para(doc, "The final implementation demonstrates that combining visual learning with landmark geometry can improve sign recognition compared to relying on raw images alone. The project also shows the importance of matching training preprocessing with live prediction preprocessing. By using landmarks and tight crops in both stages, the system becomes more consistent. The no-hand mechanism completes the interaction loop by allowing the user to spell letters and trigger speech output naturally by removing the hand from the frame.")
    add_para(doc, "Overall, the project provides a useful foundation for assistive communication. With additional data collection, stronger evaluation across multiple users, and improved speech recognition, the system can become more accurate, accessible, and reliable for real-world communication support.")

    add_heading(doc, "REFERENCES", 1)
    refs = [
        "Graves, A., Fernandez, S., Gomez, F., & Schmidhuber, J. (2006). Connectionist temporal classification: Labelling unsegmented sequence data with recurrent neural networks. Proceedings of the International Conference on Machine Learning.",
        "Lugaresi, C., Tang, J., Nash, H., McClanahan, C., Uboweja, E., Hays, M., Zhang, F., Chang, C., Yong, M., Lee, J., Chang, W., Hua, W., Georg, M., & Grundmann, M. (2019). MediaPipe: A framework for building perception pipelines. Google Research.",
        "Paszke, A., Gross, S., Massa, F., Lerer, A., Bradbury, J., Chanan, G., Killeen, T., Lin, Z., Gimelshein, N., Antiga, L., Desmaison, A., Kopf, A., Yang, E., DeVito, Z., Raison, M., Tejani, A., Chilamkurthy, S., Steiner, B., Fang, L., Bai, J., & Chintala, S. (2019). PyTorch: An imperative style, high-performance deep learning library. Advances in Neural Information Processing Systems.",
        "Bradski, G. (2000). The OpenCV library. Dr. Dobb's Journal of Software Tools.",
        "Russell, S., & Norvig, P. (2021). Artificial Intelligence: A Modern Approach. Pearson.",
        "World Health Organization. (2021). World report on hearing. World Health Organization.",
        "Doe, J. (2022). Deep learning approaches for sign language recognition. Journal of Assistive AI Systems, 4(2), 44-58. [Dummy reference for formatting].",
        "Smith, A., & Uwimana, P. (2023). Real-time gesture translation using computer vision. International Journal of Human-Computer Accessibility, 7(1), 15-31. [Dummy reference for formatting].",
        "Okafor, C. (2024). Hybrid landmark and convolutional models for alphabet-level fingerspelling. African Journal of Applied Machine Learning, 2(3), 101-118. [Dummy reference for formatting].",
    ]
    for ref in refs:
        add_para(doc, ref)

    add_heading(doc, "APPENDICES", 1)
    add_heading(doc, "Appendix A: Proposed System Requirements", 2)
    add_table(doc, "Table A.1: Hardware and Software Requirements", ["Requirement Type", "Minimum Requirement", "Recommended Requirement"], [
        ["Processor", "Dual-core CPU", "Modern multi-core CPU or CUDA-capable GPU"],
        ["Memory", "8 GB RAM", "16 GB RAM or higher"],
        ["Camera", "Built-in webcam", "External HD webcam with stable lighting"],
        ["Python", "Python 3.10", "Python 3.10 virtual environment"],
        ["Libraries", "PyTorch, OpenCV, MediaPipe, NumPy, SpeechRecognition, pyttsx3", "Same libraries with GPU-enabled PyTorch where available"],
    ])
    add_heading(doc, "Appendix B: Suggested Future Dataset Recording Plan", 2)
    add_para(doc, "For future improvement, each alphabet letter should be recorded under controlled but varied conditions. At least 20 to 50 short clips per letter should be collected from different users. Each recording should contain a clear stable pose and small natural movement. The recording tool should save the raw frame, tight hand crop, landmark vector, label, signer ID, date, and lighting condition.")
    add_fig(doc, "B.1", "Dataset recording interface with target letter prompt and webcam preview")
    add_table(doc, "Table B.1: Sample Data Collection Schedule", ["Session", "Letters", "Samples per Letter", "Notes"], [
        ["Session 1", "A-F", "30 clips and 50 images", "Focus on closed-hand letters and thumb placement."],
        ["Session 2", "G-L", "30 clips and 50 images", "Focus on sideways hand orientation and index finger signs."],
        ["Session 3", "M-R", "30 clips and 50 images", "Focus on thumb-under-finger and curved-hand letters."],
        ["Session 4", "S-Z", "30 clips and 50 images", "Focus on visually similar letters and open-finger signs."],
    ])
    add_heading(doc, "Appendix C: Example User Operation Flow", 2)
    add_numbered(doc, [
        "Open the application using python realtime/inference_ctc.py --model checkpoints/best_model.pt --no-hand-frames 8.",
        "Select sign-to-speech mode if the user wants to spell using hand signs.",
        "Place the hand inside the camera frame and form one ASL letter clearly for about one second.",
        "Wait for the letter to appear in the prediction section.",
        "Continue spelling until the word is complete.",
        "Remove the hand from the frame to complete and speak the word.",
        "Switch to speech-to-sign mode if the user wants to speak and view sign images.",
        "Select the correct microphone index and speak clearly.",
        "View the recognized text and corresponding ASL letter images.",
    ])
    add_heading(doc, "Appendix D: Image Placement Checklist", 2)
    add_para(doc, "The following image placeholders were intentionally inserted in the report and should be replaced with screenshots, diagrams, or dataset examples before final submission:")
    add_list(doc, [
        "Figure 3.1 - Dataset folder structure.",
        "Figure 3.2 - MediaPipe landmarks on webcam frame.",
        "Figure 3.3 - Hybrid training flow.",
        "Figure 3.4 - One-second prediction window.",
        "Figure 3.5 - No-hand detection flow.",
        "Figure 3.6 - Threaded runtime architecture.",
        "Figure 4.1 - Overall system architecture.",
        "Figure 4.4 - Model architecture.",
        "Figure 4.6 - Sign-to-speech interface.",
        "Figure 4.8 - Speech-to-sign interface.",
        "Figure 5.2 - Training result screenshot.",
        "Figure 5.5 - Future data collection screen.",
    ])

    doc.save(OUT)
    print(OUT)
    print(f"paragraphs={len(doc.paragraphs)}")
    print(f"tables={len(doc.tables)}")


if __name__ == "__main__":
    build()
